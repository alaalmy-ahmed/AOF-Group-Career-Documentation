"""
Generate a fake operations workbook with the same structure as the real one.

Everything here is made up: managers, branch codes, trainees, employee IDs
and issues. The workbook mirrors the production file's sheets and columns, and
stores the issue photos the same way Excel does for "Place in Cell" pictures
(rich values), so the daily notebook's image extraction works on it unchanged.

Usage:
    python sample_data/generate_sample_data.py

Outputs (next to this script):
    sample_workbook.xlsx   Source Data, Reviewing Reports, branches and codes,
                           Managers, Fixed Choices
    managers.xlsx          Wide format: one column per manager, branches below
"""

import io
import os
import random
import re
import zipfile
from datetime import date, timedelta

import openpyxl
import pandas as pd
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBOOK_PATH = os.path.join(HERE, "sample_workbook.xlsx")
MANAGERS_PATH = os.path.join(HERE, "managers.xlsx")

SEED = 7
END_DATE = date(2026, 10, 26)   # last day of data
DAYS = 14

# --------------------------------------------------------------------------------------------------------
# Reference data (all fictional)

MANAGERS = {
    "North Riyadh Manager":   ["Al Nakheel - النخيل", "Al Yasmin - الياسمين", "Al Malqa - الملقا", "Al Narjis - النرجس"],
    "East Riyadh Manager":    ["Al Rawdah - الروضة", "Al Nahdah - النهضة", "Al Qadisiyah - القادسية", "Al Rimal - الرمال"],
    "West Riyadh Manager":    ["Al Suwaidi - السويدي", "Al Badiah - البديعة", "Laban - لبن", "Dhahrat Laban - ظهرة لبن"],
    "South Riyadh Manager":   ["Al Aziziyah - العزيزية", "Al Shifa - الشفا", "Al Dar Al Baida - الدار البيضاء"],
    "Central Riyadh Manager": ["Al Olaya - العليا", "Al Malaz - الملز", "Al Murabba - المربع", "Al Wurud - الورود"],
    "Outskirts Manager":      ["Al Kharj - الخرج", "Al Diriyah - الدرعية", "Al Muzahimiyah - المزاحمية"],
    "Jeddah North Manager":   ["Obhur - أبحر", "Al Hamdaniyah - الحمدانية", "Al Marwah - المروة", "Al Safa - الصفا"],
    "Jeddah South Manager":   ["Al Faisaliyah - الفيصلية", "Al Rabwah - الربوة", "Al Naseem - النسيم"],
    "Dammam Manager":         ["Al Faisaliyah Dammam - الفيصلية الدمام", "Al Shati - الشاطئ", "Al Aziziyah Khobar - العزيزية الخبر", "Al Rakah - الراكة"],
    "Al Ahsa Manager":        ["Al Hofuf - الهفوف", "Al Mubarraz - المبرز", "Al Rawdah Hofuf - الروضة الهفوف"],
}

MEALS = [
    ("Shawarma Classic", "8-12 pm"),
    ("Arabi Sop", "12-4 pm"),
    ("Heating Shawarma", "8-12 pm"),
    ("Cutting Shawarma", "4-8 pm"),
    ("French Fries", "12-4 pm"),
]

ISSUES = [
    "Wrong sauce quantity",
    "Bread not heated to standard",
    "Portion size below standard",
    "Wrapping not as per SOP",
    "Missing tools at the station",
    "Cutting thickness not as per SOP",
    "Fries over-fried",
    "Holding time exceeded",
    "Station not cleaned between orders",
    "Garnish missing",
]

TRAINERS = ["Trainer A", "Trainer B", "Trainer C", "Trainer D"]
FIRST_NAMES = ["Khalid", "Faisal", "Omar", "Yousef", "Hassan", "Saeed", "Nasser", "Bilal",
               "Tariq", "Majed", "Ali", "Rakan", "Imran", "Sami", "Adel", "Ziyad"]
LAST_INITIALS = list("ABDFHKMNRSTZ")

PHOTO_COLORS = ["#C8553D", "#F28F3B", "#588B8B", "#2D3047", "#93B7BE", "#8C5383"]


# --------------------------------------------------------------------------------------------------------
# Data generation

def build_frames(rng):
    branches, managers_long = [], []
    for i, (manager, names) in enumerate(MANAGERS.items()):
        for name in names:
            code = f"FAKE B{len(branches) + 1:02d}"
            branches.append({"Code": code, "Branch": name})
            managers_long.append({"Branch": name, "Manager": manager})
    branches_df = pd.DataFrame(branches)
    managers_df = pd.DataFrame(managers_long)
    branch_to_manager = dict(zip(managers_df["Branch"], managers_df["Manager"]))
    branch_to_code = dict(zip(branches_df["Branch"], branches_df["Code"]))
    all_branches = branches_df["Branch"].tolist()

    # A few branches struggle more than others, so the weekly ranking has a spread.
    weak = set(rng.sample(all_branches, 6))
    weights = [5 if b in weak else 1 for b in all_branches]

    trainees = [(f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_INITIALS)}.", rng.randint(1000000, 1999999))
                for _ in range(60)]

    tasks, issues = [], []
    for d in range(DAYS):
        day = END_DATE - timedelta(days=DAYS - 1 - d)
        for meal, due in MEALS:
            n_missed = rng.choice([0, 0, 1, 1, 2, 3, 4])
            missed = sorted(set(rng.choices(all_branches, weights=weights, k=n_missed)))
            done = len(all_branches) - len(missed)
            tasks.append({
                "Date": pd.Timestamp(day),
                "Meal": meal,
                "Team": "S.C Branches",
                "Done Tasks": done,
                "Missed Tasks": len(missed),
                "Branches Missed Tasks": "\n".join(missed) if missed else None,
                "Due time": due,
                "Team Done Rate": round(done / len(all_branches), 3),
            })

        for _ in range(rng.randint(4, 9)):
            branch = rng.choices(all_branches, weights=weights, k=1)[0]
            trainee, emp_id = rng.choice(trainees)
            issues.append({
                "Date": pd.Timestamp(day),
                "Code": branch_to_code[branch],
                "Branch": branch,
                "Prepared By": rng.choice(TRAINERS),
                "Trainee Name": trainee,
                "Employee ID": emp_id,
                "Meal": rng.choice(MEALS)[0],
                "Description": rng.choice(ISSUES),
                "Manager": branch_to_manager[branch],
                "Images": "#VALUE!",   # placeholder; becomes an in-cell picture below
            })

    tasks_df = pd.DataFrame(tasks)
    issues_df = pd.DataFrame(issues)
    fixed_df = pd.DataFrame({"Issues": ISSUES, "Y/N": ["Y"] * len(ISSUES)})
    return tasks_df, issues_df, branches_df, managers_df, fixed_df


def make_photo(i):
    """Small placeholder 'photo' so the output files show where real images go."""
    img = Image.new("RGB", (400, 400), PHOTO_COLORS[i % len(PHOTO_COLORS)])
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, 379, 379], outline="white", width=6)
    draw.text((140, 185), f"Sample photo {i + 1}", fill="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------------------
# In-cell pictures
#
# openpyxl can't write Excel's "Place in Cell" pictures, so the workbook is
# written first and the rich-value parts are added to the zip afterwards:
#   cell vm="n" -> metadata.xml valueMetadata[n-1] -> rich value -> richValueRel -> xl/media/imageN.png

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_RD = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"
NS_RD2 = "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata2"
NS_RVR = "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def rich_parts(n_images):
    rv = "".join(f'<rv s="0"><v>{i}</v><v>5</v></rv>' for i in range(n_images))
    rels = "".join(f'<rel r:id="rId{i + 1}"/>' for i in range(n_images))
    rel_targets = "".join(
        f'<Relationship Id="rId{i + 1}" Type="{NS_R}/image" Target="../media/image{i + 1}.png"/>'
        for i in range(n_images))
    future = "".join(
        f'<bk><extLst><ext uri="{{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}}"><xlrd:rvb i="{i}"/></ext></extLst></bk>'
        for i in range(n_images))
    value_meta = "".join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(n_images))
    return {
        "xl/richData/rdRichValueTypes.xml": XML_HEAD +
            f'<rvTypesInfo xmlns="{NS_RD2}"><global><keyFlags>'
            '<key name="_Self"><flag name="ExcludeFromFile" value="1"/><flag name="ExcludeFromCalcComparison" value="1"/></key>'
            '<key name="_DisplayString"><flag name="ExcludeFromCalcComparison" value="1"/></key>'
            '</keyFlags></global></rvTypesInfo>',
        "xl/richData/rdrichvaluestructure.xml": XML_HEAD +
            f'<rvStructures xmlns="{NS_RD}" count="1"><s t="_localImage">'
            '<k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>',
        "xl/richData/rdrichvalue.xml": XML_HEAD + f'<rvData xmlns="{NS_RD}" count="{n_images}">{rv}</rvData>',
        "xl/richData/richValueRel.xml": XML_HEAD + f'<richValueRels xmlns="{NS_RVR}" xmlns:r="{NS_R}">{rels}</richValueRels>',
        "xl/richData/_rels/richValueRel.xml.rels": XML_HEAD + f'<Relationships xmlns="{NS_PKG}">{rel_targets}</Relationships>',
        "xl/metadata.xml": XML_HEAD +
            f'<metadata xmlns="{NS_MAIN}" xmlns:xlrd="{NS_RD}">'
            '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000" copy="1" '
            'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
            'clearComments="1" assign="1" coerce="1"/></metadataTypes>'
            f'<futureMetadata name="XLRICHVALUE" count="{n_images}">{future}</futureMetadata>'
            f'<valueMetadata count="{n_images}">{value_meta}</valueMetadata></metadata>',
    }


def add_in_cell_pictures(xlsx_path, sheet_xml, image_col, row_to_vm, photos):
    with zipfile.ZipFile(xlsx_path) as z:
        parts = {name: z.read(name) for name in z.namelist()}

    # Tag each picture cell with its vm index.
    sheet = parts[sheet_xml].decode()
    def tag(m):
        row = int(m.group(2))
        vm = row_to_vm.get(row)
        return m.group(0) if vm is None else f'{m.group(1)} vm="{vm}"'
    sheet = re.sub(rf'(<c r="{image_col}(\d+)"[^>]*?)(?=\s*/?>)', tag, sheet)
    parts[sheet_xml] = sheet.encode()

    for name, xml in rich_parts(len(photos)).items():
        parts[name] = xml.encode()
    for i, png in enumerate(photos):
        parts[f"xl/media/image{i + 1}.png"] = png

    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    extra = (
        f'<Relationship Id="rId101" Type="{NS_R}/sheetMetadata" Target="metadata.xml"/>'
        '<Relationship Id="rId102" Type="http://schemas.microsoft.com/office/2022/10/relationships/richValueRel" Target="richData/richValueRel.xml"/>'
        '<Relationship Id="rId103" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue" Target="richData/rdrichvalue.xml"/>'
        '<Relationship Id="rId104" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure" Target="richData/rdrichvaluestructure.xml"/>'
        '<Relationship Id="rId105" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueTypes" Target="richData/rdRichValueTypes.xml"/>'
    )
    parts["xl/_rels/workbook.xml.rels"] = rels.replace("</Relationships>", extra + "</Relationships>").encode()

    ct = parts["[Content_Types].xml"].decode()
    overrides = (
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>'
        '<Override PartName="/xl/richData/richValueRel.xml" ContentType="application/vnd.ms-excel.richvaluerel+xml"/>'
        '<Override PartName="/xl/richData/rdrichvalue.xml" ContentType="application/vnd.ms-excel.rdrichvalue+xml"/>'
        '<Override PartName="/xl/richData/rdrichvaluestructure.xml" ContentType="application/vnd.ms-excel.rdrichvaluestructure+xml"/>'
        '<Override PartName="/xl/richData/rdRichValueTypes.xml" ContentType="application/vnd.ms-excel.rdrichvaluetypes+xml"/>'
    )
    parts["[Content_Types].xml"] = ct.replace("</Types>", overrides + "</Types>").encode()

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as z:
        # [Content_Types].xml must come first in the package.
        z.writestr("[Content_Types].xml", parts.pop("[Content_Types].xml"))
        for name, data in parts.items():
            z.writestr(name, data)


# --------------------------------------------------------------------------------------------------------
# Main

def main():
    rng = random.Random(SEED)
    tasks_df, issues_df, branches_df, managers_df, fixed_df = build_frames(rng)

    with pd.ExcelWriter(WORKBOOK_PATH, engine="openpyxl") as writer:
        tasks_df.to_excel(writer, sheet_name="Source Data", index=False)
        issues_df.to_excel(writer, sheet_name="Reviewing Reports", index=False)
        branches_df.to_excel(writer, sheet_name="branches and codes", index=False)
        managers_df.to_excel(writer, sheet_name="Managers", index=False)
        fixed_df.to_excel(writer, sheet_name="Fixed Choices", index=False)

    # Write the Images cells as #VALUE! errors, which is how Excel stores an in-cell picture's value.
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["Reviewing Reports"]
    image_col = openpyxl.utils.get_column_letter(list(issues_df.columns).index("Images") + 1)
    for row in range(2, len(issues_df) + 2):
        ws[f"{image_col}{row}"].value = "#VALUE!"
        ws[f"{image_col}{row}"].data_type = "e"
    sheet_index = wb.sheetnames.index("Reviewing Reports") + 1
    wb.save(WORKBOOK_PATH)

    photos = [make_photo(i) for i in range(len(PHOTO_COLORS))]
    row_to_vm = {row: rng.randint(1, len(photos)) for row in range(2, len(issues_df) + 2)}
    add_in_cell_pictures(WORKBOOK_PATH, f"xl/worksheets/sheet{sheet_index}.xml", image_col, row_to_vm, photos)

    # Wide managers file used by the daily notebook.
    wide = pd.DataFrame({m: pd.Series(b) for m, b in MANAGERS.items()})
    wide.to_excel(MANAGERS_PATH, index=False)

    print(f"Wrote {WORKBOOK_PATH}")
    print(f"  {len(tasks_df)} task rows, {len(issues_df)} issues, {len(branches_df)} branches, "
          f"{len(MANAGERS)} managers, {len(photos)} photos, {tasks_df['Date'].min().date()} to {END_DATE}")
    print(f"Wrote {MANAGERS_PATH}")


if __name__ == "__main__":
    main()
